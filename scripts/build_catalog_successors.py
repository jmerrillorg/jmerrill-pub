"""Build source-backed successor catalog drafts from the live Dataverse catalog."""

import argparse
import csv
import json
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from catalog_authority_audit import REGISTER, live_rows


STATUS = {
    100000000: "Active",
    100000001: "Superseded",
    100000002: "Retired",
    100000003: "Provisional",
    100000004: "Internal only",
}


def dollars(value):
    return f"${float(value):,.2f}" if value is not None else "Price by approved quote or scope"


def base_document(title, subtitle):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    document.styles["Title"].font.name = "Aptos Display"
    document.styles["Title"].font.size = Pt(20)
    document.styles["Title"].font.color.rgb = None
    title_style = document.styles["Title"]
    style_properties = title_style.element.get_or_add_pPr()
    border = style_properties.find(qn("w:pBdr"))
    if border is not None:
        style_properties.remove(border)
    title_paragraph = document.add_paragraph(title, "Title")
    paragraph_properties = title_paragraph._p.get_or_add_pPr()
    border = paragraph_properties.find(qn("w:pBdr"))
    if border is not None:
        paragraph_properties.remove(border)
    document.add_paragraph(subtitle)
    return document


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def build_full_catalog(rows, output):
    document = base_document("Full Products and Services Catalog", "J Merrill Publishing  |  Successor review draft  |  1 October 2026")
    document.add_paragraph("J Merrill Publishing has published 130+ titles. This is the founder-confirmed portfolio statement, not a count of records visible in any particular live catalog view. This draft presents the current operational service register without promoting unresolved commercial terms.")
    document.add_heading("Publishing packages", level=1)
    document.add_paragraph("The three standard packages are Starter, Professional, and Premier. The historical Signature package is no longer available for new quotes or agreements. Existing signed Signature agreements remain governed by their own terms.")
    packages = [row for row in rows if row["jm1pub_canonicalsku"] in {"JMP-PKG-STARTER", "JMP-PKG-PRO", "JMP-PKG-PREMIER"}]
    add_table(document, ["Package", "SKU", "Base price"], [(row["jm1pub_name"], row["jm1pub_canonicalsku"], dollars(row.get("jm1pub_unitprice"))) for row in sorted(packages, key=lambda row: row["jm1pub_unitprice"])])
    document.add_paragraph("Children's Publishing is a separate specialty-package record at $2,495 in the August approved register. Its continuing availability under the later three-tier standard lineup awaits an explicit commercial ruling; this draft does not offer it as a fourth standard tier.")
    document.add_heading("Payment terms", level=1)
    document.add_paragraph("Flexible payment options may be available. Publishing projects may qualify for structured payment arrangements based on package, project scope, and approved payment terms. Final payment terms are provided with the publishing agreement.")
    document.add_heading("Current services", level=1)
    document.add_paragraph("Only active, public, sellable service records appear below. A missing fixed price means a governed quote or scope determination is required; it does not imply a zero-cost service. Provisional, retired, superseded, and internal-only records are excluded.")
    groups = defaultdict(list)
    with REGISTER.open(newline="") as source:
        categories = {row["Row ID"]: row["Original Category"] for row in csv.DictReader(source)}
    for row in rows:
        if row["jm1pub_commercialstatus"] == 100000000 and row["jm1pub_sellablestatus"] == 100000000 and row["jm1pub_publicvisibility"] == 100000000 and not row["jm1pub_canonicalsku"].startswith("JMP-PKG-"):
            groups[categories.get(row["jm1pub_catalogrowid"], "Specialist production")].append(row)
    for category, items in sorted(groups.items()):
        document.add_heading(category, level=2)
        add_table(document, ["Service", "SKU", "Price"], [(row["jm1pub_name"], row["jm1pub_canonicalsku"], dollars(row.get("jm1pub_unitprice"))) for row in sorted(items, key=lambda row: row["jm1pub_name"])])
    document.add_paragraph("This successor draft is not a contract, payment schedule, or approval of provisional services. Specific scope, tax, and payment terms require the applicable executed agreement and governed financial authority.")
    document.save(output)


def build_reference_guide(rows, output):
    document = base_document("Product Reference Guide", "J Merrill Publishing  |  Internal successor review draft  |  1 October 2026")
    document.add_paragraph("This guide reconciles the 120 live commercial-catalog records with the August approved service register and the later package addendum. It preserves the founder-confirmed 130+ published-title portfolio statement while keeping live catalog-listing counts separate. It is for operator review, not an author-facing quote or contract.")
    document.add_heading("Current authority", level=1)
    document.add_paragraph("Starter $1,999, Professional $4,500, and Premier $7,500 are the current standard package base prices. Premier replaces Signature for new business; historical Signature agreements are not rewritten. The Children's specialty package remains an explicit decision pending reaffirmation.")
    document.add_heading("Payment rules", level=1)
    document.add_paragraph("The versioned internal calculator supports Full Pay and 2, 4, 8, 12, 18, and 24 payments. Financing v1.1 uses a 6% annual simple plan charge prorated to financed months, without compounding; early payoff waives unearned future charges. Existing legacy contracts may have a separately versioned 4% transaction-fee model. Payment-policy version and contracted principal must be bound per agreement. No Full Pay discount is approved. Tax is external to the calculator, and public presentation remains summary only. Do not calculate a universal public schedule from these rules.")
    document.add_heading("Commercial status and source", level=1)
    document.add_paragraph("The August 5 SKU ruling register governs services where no later explicit authority changes them. The August 12 package addendum and October 1 founder correction override its older Signature-package treatment. The two merged partner-program rows CAT-111 and CAT-112 are intentionally absent from live catalog records; CAT-121 Premier and the separately approved PF08-AUTH-001 interactive EPUB3 edition are present. CAT-052 is a superseded legacy alias for PF08-AUTH-001.")
    document.add_heading("Live SKU register", level=1)
    add_table(document, ["Row", "SKU", "Status", "Price", "Commercial gate"], [
        (row["jm1pub_catalogrowid"], row["jm1pub_canonicalsku"], STATUS[row["jm1pub_commercialstatus"]], dollars(row.get("jm1pub_unitprice")), "Quote and contract" if row["jm1pub_quotingstatus"] == 100000000 and row["jm1pub_contractstatus"] == 100000000 else "Restricted")
        for row in sorted(rows, key=lambda row: row["jm1pub_catalogrowid"])
    ])
    document.add_heading("Outstanding authority", level=1)
    document.add_paragraph("Confirm whether JMP-PKG-CHILD remains a separately quotable specialty offer, is superseded, or must be held. Do not resume QBO product updates from this review draft. Fixed-price exceptions, tax, late-payment terms, and final-delivery conditions must come from the governing agreement or policy, not this document.")
    document.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    rows = live_rows()
    if len(rows) != 120:
        raise RuntimeError(f"Expected 120 live rows, found {len(rows)}")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    full = args.output_dir / "JMP_Full_Products_and_Services_Catalog_v3_REVIEW_DRAFT.docx"
    guide = args.output_dir / "JMP_Product_Reference_Guide_v2_REVIEW_DRAFT.docx"
    build_full_catalog(rows, full)
    build_reference_guide(rows, guide)
    print(json.dumps({"full_catalog": str(full), "reference_guide": str(guide), "live_rows": len(rows), "review_draft": True}, indent=2))


if __name__ == "__main__":
    main()
