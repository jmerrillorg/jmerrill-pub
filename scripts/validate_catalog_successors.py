"""Validate generated successor documents against the current live catalog."""

import argparse
import json
from pathlib import Path

from docx import Document

from build_catalog_successors import STATUS, dollars
from catalog_authority_audit import live_rows


def table_rows(document):
    return [tuple(cell.text for cell in row.cells) for table in document.tables for row in table.rows]


def all_text(document):
    return "\n".join([paragraph.text for paragraph in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])


def validate(output_dir):
    rows = live_rows()
    full = Document(output_dir / "JMP_Full_Products_and_Services_Catalog_v3.docx")
    guide = Document(output_dir / "JMP_Product_Reference_Guide_v2.docx")
    full_rows = table_rows(full)
    guide_rows = table_rows(guide)
    full_text = all_text(full)
    guide_text = all_text(guide)
    assert len(rows) == 120
    assert len({row["jm1pub_canonicalsku"] for row in rows}) == 120
    for row in rows:
        expected = (
            row["jm1pub_catalogrowid"],
            row["jm1pub_canonicalsku"],
            STATUS[row["jm1pub_commercialstatus"]],
            dollars(row.get("jm1pub_unitprice")),
            "Quote and contract" if row["jm1pub_quotingstatus"] == 100000000 and row["jm1pub_contractstatus"] == 100000000 else "Restricted",
        )
        assert guide_rows.count(expected) == 1, expected
        sku = row["jm1pub_canonicalsku"]
        if row["jm1pub_commercialstatus"] == 100000000 and row["jm1pub_sellablestatus"] == 100000000 and row["jm1pub_publicvisibility"] == 100000000:
            assert sum(sku in cells for cells in full_rows) == 1, sku
        else:
            assert not any(sku in cells for cells in full_rows), sku
    assert ("Children's Book Publishing Package", "JMP-PKG-CHILD", "$2,495.00") in full_rows
    assert "author supplies production-usable" in full_text
    assert "130+ titles" in full_text and "130+ published-title" in guide_text
    assert "Flexible payment options may be available" in full_text
    assert "Uncategorized" not in full_text
    for text in (full_text, guide_text):
        for stale in ("7%", "6% annual simple", "pending Children's", "review draft"):
            assert stale not in text, stale
    assert "summary only" in guide_text
    return {"live_rows": len(rows), "active_rows": sum(row["jm1pub_commercialstatus"] == 100000000 for row in rows), "full_sku_rows": sum(any(row["jm1pub_canonicalsku"] in cells for cells in full_rows) for row in rows), "guide_sku_rows": len(rows), "parity": "PASS"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(validate(args.output_dir), indent=2))
